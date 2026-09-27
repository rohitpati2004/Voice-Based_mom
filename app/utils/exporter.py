import os
from pathlib import Path
from typing import Dict, Any
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
import docx

class DocumentExporter:
    @staticmethod
    def export_pdf(mom_data: Dict[str, Any], output_path: str | Path) -> str:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#1E293B'),
            spaceAfter=12
        )
        heading_style = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor('#0F172A'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#334155')
        )
        bold_body = ParagraphStyle(
            'BoldBody',
            parent=body_style,
            fontName='Helvetica-Bold'
        )

        elements = []

        # Document Header
        title_text = mom_data.get("meeting_title", "Voice-Based Minutes of Meeting")
        elements.append(Paragraph(title_text, title_style))
        elements.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#3B82F6'), spaceAfter=15))

        # Executive Summary
        elements.append(Paragraph("Executive Summary", heading_style))
        elements.append(Paragraph(mom_data.get("summary", "N/A"), body_style))
        elements.append(Spacer(1, 10))

        # Speaker Statistics Table
        stats = mom_data.get("statistics", {}).get("speaker_statistics", [])
        if stats:
            elements.append(Paragraph("Speaker Participation Statistics", heading_style))
            table_data = [["Speaker", "Speaking Time", "Proportion", "Segments", "Language"]]
            for s in stats:
                table_data.append([
                    s["speaker"],
                    s["speaking_duration_formatted"],
                    f"{s['speaking_proportion_percent']}%",
                    str(s["segment_count"]),
                    s["primary_language"]
                ])
            
            t = Table(table_data, colWidths=[100, 100, 100, 80, 120])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F8FAFC')),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#334155')),
            ]))
            elements.append(t)
            elements.append(Spacer(1, 12))

        # Key Discussion Points
        key_points = mom_data.get("key_discussion_points", [])
        if key_points:
            elements.append(Paragraph("Key Discussion Points", heading_style))
            for kp in key_points:
                elements.append(Paragraph(f"<b>• {kp['topic']}</b>", body_style))
                for pt in kp.get("points", []):
                    elements.append(Paragraph(f"&nbsp;&nbsp;&nbsp;&nbsp;- {pt}", body_style))
            elements.append(Spacer(1, 10))

        # Decisions Made
        decisions = mom_data.get("decisions", [])
        if decisions:
            elements.append(Paragraph("Decisions Made", heading_style))
            for d in decisions:
                spk = d.get("speaker", "General")
                ts = d.get("timestamp", "")
                elements.append(Paragraph(f"• <b>[{ts}] {spk}:</b> {d['decision']}", body_style))
            elements.append(Spacer(1, 10))

        # Action Items
        actions = mom_data.get("action_items", [])
        if actions:
            elements.append(Paragraph("Action Items & Responsibilities", heading_style))
            action_table = [["Assignee", "Task", "Status", "Timestamp"]]
            for act in actions:
                action_table.append([
                    act.get("assignee", "Unassigned"),
                    act.get("task", ""),
                    act.get("status", "Pending"),
                    act.get("timestamp", "")
                ])
            at = Table(action_table, colWidths=[100, 240, 80, 80])
            at.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F1F5F9')),
            ]))
            elements.append(at)
            elements.append(Spacer(1, 12))

        # Full Speaker Transcript
        transcript = mom_data.get("transcript", [])
        if transcript:
            elements.append(Paragraph("Speaker-Wise Chronological Transcript", heading_style))
            for seg in transcript:
                start_m = int(seg['start'] // 60)
                start_s = int(seg['start'] % 60)
                end_m = int(seg['end'] // 60)
                end_s = int(seg['end'] % 60)
                time_str = f"[{start_m:02d}:{start_s:02d} - {end_m:02d}:{end_s:02d}]"
                lang_str = f"({seg.get('language', 'En')})"
                text_line = f"<b>{time_str} {seg['speaker']} {lang_str}:</b> {seg['text']}"
                elements.append(Paragraph(text_line, body_style))

        doc.build(elements)
        return str(output_path)

    @staticmethod
    def export_docx(mom_data: Dict[str, Any], output_path: str | Path) -> str:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()
        doc.add_heading(mom_data.get("meeting_title", "Voice-Based Minutes of Meeting"), level=0)

        # Executive Summary
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(mom_data.get("summary", "N/A"))

        # Speaker Statistics
        stats = mom_data.get("statistics", {}).get("speaker_statistics", [])
        if stats:
            doc.add_heading("Speaker Participation Statistics", level=1)
            t = doc.add_table(rows=1, cols=5)
            t.style = 'Table Grid'
            hdr = t.rows[0].cells
            hdr[0].text = "Speaker"
            hdr[1].text = "Speaking Duration"
            hdr[2].text = "Proportion"
            hdr[3].text = "Segments"
            hdr[4].text = "Primary Language"
            for s in stats:
                row = t.add_row().cells
                row[0].text = s["speaker"]
                row[1].text = s["speaking_duration_formatted"]
                row[2].text = f"{s['speaking_proportion_percent']}%"
                row[3].text = str(s["segment_count"])
                row[4].text = s["primary_language"]

        # Key Discussion Points
        key_points = mom_data.get("key_discussion_points", [])
        if key_points:
            doc.add_heading("Key Discussion Points", level=1)
            for kp in key_points:
                p = doc.add_paragraph(style='List Bullet')
                r = p.add_run(kp['topic'])
                r.bold = True
                for pt in kp.get("points", []):
                    doc.add_paragraph(pt, style='List Bullet 2')

        # Decisions Made
        decisions = mom_data.get("decisions", [])
        if decisions:
            doc.add_heading("Decisions Made", level=1)
            for d in decisions:
                spk = d.get("speaker", "General")
                ts = d.get("timestamp", "")
                doc.add_paragraph(f"[{ts}] {spk}: {d['decision']}", style='List Bullet')

        # Action Items
        actions = mom_data.get("action_items", [])
        if actions:
            doc.add_heading("Action Items & Responsibilities", level=1)
            at = doc.add_table(rows=1, cols=4)
            at.style = 'Table Grid'
            ahdr = at.rows[0].cells
            ahdr[0].text = "Assignee"
            ahdr[1].text = "Task"
            ahdr[2].text = "Status"
            ahdr[3].text = "Timestamp"
            for act in actions:
                arow = at.add_row().cells
                arow[0].text = act.get("assignee", "Unassigned")
                arow[1].text = act.get("task", "")
                arow[2].text = act.get("status", "Pending")
                arow[3].text = act.get("timestamp", "")

        # Transcript
        transcript = mom_data.get("transcript", [])
        if transcript:
            doc.add_heading("Speaker-Wise Chronological Transcript", level=1)
            for seg in transcript:
                start_m = int(seg['start'] // 60)
                start_s = int(seg['start'] % 60)
                end_m = int(seg['end'] // 60)
                end_s = int(seg['end'] % 60)
                t_str = f"[{start_m:02d}:{start_s:02d} - {end_m:02d}:{end_s:02d}]"
                doc.add_paragraph(f"{t_str} {seg['speaker']} ({seg.get('language', 'En')}): {seg['text']}")

        doc.save(str(output_path))
        return str(output_path)
